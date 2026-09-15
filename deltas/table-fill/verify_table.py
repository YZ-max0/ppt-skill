#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""读回验证（表格）：递归统计 Table shape / 行列数 / 单元格文本抽样。

**关键点（沿用 T-IMG1 C-029 教训）**：本仓库骨架的根级 `<g>` 会成组，
所有 shape 枚举**必须递归进 group**，否则会漏。

判据：
  · `shape.has_table == True` → 原生 PowerPoint 表格对象（可编辑行列）
  · 否则若含文本 → 仅是普通文本框（说明走的是 SVG 回退路径）
"""
from __future__ import annotations

import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def walk(shapes, depth=0):
    """递归产出 (depth, shape)。"""
    for sh in shapes:
        yield depth, sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes, depth + 1)


def main(pptx_path: str) -> int:
    prs = Presentation(pptx_path)
    print(f"slides: {len(prs.slides._sldIdLst)}")
    total_tables = 0
    for si, slide in enumerate(prs.slides, 1):
        tables, textboxes = [], 0
        for depth, sh in walk(slide.shapes):
            try:
                if sh.has_table:
                    tables.append((depth, sh))
                    continue
            except Exception:
                pass
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                textboxes += 1
        print(f"\n--- slide {si}: {len(tables)} table shape(s), "
              f"{textboxes} 非空文本框 ---")
        for depth, sh in tables:
            total_tables += 1
            tb = sh.table
            rows, cols = len(tb.rows), len(tb.columns)
            w_px, h_px = sh.width / 9525, sh.height / 9525
            print(f"  [表] depth={depth} {sh.name!r}  {rows} 行 x {cols} 列  "
                  f"显示 {w_px:.0f}x{h_px:.0f}px")
            # 单元格文本抽查：表头行 + 第一数据行
            for r in range(min(rows, 2)):
                cells = [tb.cell(r, c).text.strip() for c in range(cols)]
                print(f"       row{r}: {cells}")
        # 无原生表时，提示走的是回退路径
        if not tables and textboxes:
            print("       (无原生 Table；该页表格以矢量形状/文本框呈现)")
    print(f"\nTOTAL native table shapes = {total_tables}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1]))
